from openpyxl.chart import PieChart, Reference
from funciones import *

ruta = '../Data/data.xlsx'

archivo_excel = cargar_archivo( ruta )
hoja = obtener_hoja(archivo_excel, 'Seguidores' )

grafico = PieChart()

categorias = Reference( hoja, min_col=2, min_row=2, max_row=hoja.max_row )
datos = Reference( hoja, min_col=3, min_row=2, max_row=hoja.max_row )

grafico.add_data( datos )
grafico.set_categories( categorias )
grafico.title = 'Seguidores'

hoja.add_chart( grafico, 'E2' )

archivo_excel.save( ruta )